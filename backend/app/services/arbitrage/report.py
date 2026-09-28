"""
report.py — Xuất báo cáo XLSX cho kết quả quét đối đả.

Output:
1. File chính: Kết quả đối đả
   - Sheet "总表" (tổng hợp)
   - Sheet riêng cho mỗi đài (chứa 内对打 + 外对打 liên quan)

2. File loại: 同手淘汰记录
   - Danh sách cặp bị loại do cùng tay ≥2

Quy tắc:
- Tài khoản lưu dạng text (giữ số 0 đầu)
- Tiền hiển thị 1,234.00
- 外对打 xuất ở cả 2 sheet liên quan
- Trong sheet đài nào → tài khoản đài đó = 本站账号
- 对打组 = tên đài hiện tại đứng trước
"""

from __future__ import annotations

import os
from typing import BinaryIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, numbers
from openpyxl.utils import get_column_letter

from .detector import ArbitragePair, EliminatedPair, DetectionResult
from .rules import display_platform


# ---------------------------------------------------------------------------
# Style constants
# ---------------------------------------------------------------------------

_HEADER_FONT = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

_DATA_FONT = Font(name="Microsoft YaHei", size=10)
_DATA_ALIGN = Alignment(horizontal="center", vertical="center")
_TEXT_ALIGN = Alignment(horizontal="left", vertical="center")

_MONEY_FORMAT = '#,##0.00'
_PCT_FORMAT = '0.00%'

_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

# Colors for match types
_INTERNAL_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
_EXTERNAL_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")


# ---------------------------------------------------------------------------
# Column definitions (theo đúng mục 8 & 13 trong tài liệu quy định)
# ---------------------------------------------------------------------------

_MAIN_HEADERS = [
    "对打组",        # match group: platform_a-platform_b
    "本站账号",      # local account
    "其他站账号",    # remote account
    "Ván",           # number of arb rounds
    "Bằng tiền",     # number of equal-stake rounds
    "Tổng tiền",     # total stake both sides
    "Lệch",          # total stake difference
    "备注",          # remarks
]

_ELIMINATED_HEADERS = [
    "对打组",
    "账号A",
    "账号B",
    "Ván đối打",
    "Bằng tiền",
    "Tổng tiền",
    "Lệch",
    "同手Ván",
    "备注",
]


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def generate_report(
    result: DetectionResult,
    output_dir: str,
    filename_prefix: str = "对打结果",
) -> dict[str, str]:
    """Xuất báo cáo XLSX từ DetectionResult.

    Args:
        result: Kết quả quét từ ArbitrageDetector.detect().
        output_dir: Thư mục output.
        filename_prefix: Tiền tố tên file.

    Returns:
        Dict {"main": path, "eliminated": path} — đường dẫn file đã tạo.
    """
    os.makedirs(output_dir, exist_ok=True)

    main_path = os.path.join(output_dir, f"{filename_prefix}.xlsx")
    elim_path = os.path.join(output_dir, f"{filename_prefix}_同手淘汰记录.xlsx")

    # --- File chính ---
    _generate_main_report(result, main_path)

    # --- File loại ---
    _generate_eliminated_report(result.eliminated_pairs, elim_path)

    return {"main": main_path, "eliminated": elim_path}


# ---------------------------------------------------------------------------
# Main report
# ---------------------------------------------------------------------------

def _generate_main_report(result: DetectionResult, path: str) -> None:
    wb = Workbook()

    # Sheet 总表
    ws_summary = wb.active
    ws_summary.title = "总表"
    _write_summary_sheet(ws_summary, result)

    # Sheet riêng cho mỗi đài
    for platform in result.platforms:
        disp = display_platform(platform)
        # Tên sheet giới hạn 31 ký tự
        sheet_name = disp[:31]
        ws = wb.create_sheet(title=sheet_name)
        _write_platform_sheet(ws, platform, result)

    wb.save(path)


def _write_summary_sheet(ws, result: DetectionResult) -> None:
    """Viết sheet 总表 chứa tất cả cặp hợp lệ."""
    _write_headers(ws, _MAIN_HEADERS)

    row = 2
    for pair in result.valid_pairs:
        disp_a = display_platform(pair.platform_a)
        disp_b = display_platform(pair.platform_b)
        group_name = f"{disp_a}-{disp_b}"

        row = _write_pair_row(
            ws, row, pair, group_name,
            local_account=pair.account_a,
            remote_account=pair.account_b,
        )

    _auto_column_width(ws)


def _write_platform_sheet(
    ws, platform: str, result: DetectionResult
) -> None:
    """Viết sheet cho 1 đài: chứa 内对打 + 外对打 liên quan."""
    _write_headers(ws, _MAIN_HEADERS)

    row = 2
    for pair in result.valid_pairs:
        # Chỉ lấy pair liên quan đến đài này
        if pair.platform_a != platform and pair.platform_b != platform:
            continue

        disp_current = display_platform(platform)

        if pair.platform_a == platform and pair.platform_b == platform:
            # 内对打
            group_name = f"{disp_current}-{disp_current}"
            local = pair.account_a
            remote = pair.account_b
        elif pair.platform_a == platform:
            disp_other = display_platform(pair.platform_b)
            # Đài hiện tại đứng trước
            group_name = f"{disp_current}-{disp_other}"
            local = pair.account_a
            remote = pair.account_b
        else:
            disp_other = display_platform(pair.platform_a)
            # Đài hiện tại đứng trước
            group_name = f"{disp_current}-{disp_other}"
            # Đảo: tài khoản đài này = 本站账号
            local = pair.account_b
            remote = pair.account_a

        row = _write_pair_row(
            ws, row, pair, group_name,
            local_account=local,
            remote_account=remote,
        )

    _auto_column_width(ws)


# ---------------------------------------------------------------------------
# Eliminated report
# ---------------------------------------------------------------------------

def _generate_eliminated_report(
    eliminated: list[EliminatedPair], path: str
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "同手淘汰记录"
    _write_headers(ws, _ELIMINATED_HEADERS)

    row = 2
    for ep in eliminated:
        disp_a = display_platform(ep.platform_a)
        disp_b = display_platform(ep.platform_b)
        group_name = f"{disp_a}-{disp_b}"

        values = [
            group_name,
            ep.account_a,
            ep.account_b,
            ep.num_arb_rounds,
            ep.num_equal_stake,
            ep.total_stake,
            ep.total_diff,
            ep.num_same_hand,
            ep.remark,
        ]
        _write_data_row(ws, row, values)
        row += 1

    _auto_column_width(ws)
    wb.save(path)


# ---------------------------------------------------------------------------
# Row writers
# ---------------------------------------------------------------------------

def _write_pair_row(
    ws,
    row: int,
    pair: ArbitragePair,
    group_name: str,
    local_account: str,
    remote_account: str,
) -> int:
    """Viết 1 dòng cho 1 cặp đối đả. Return next row."""
    values = [
        group_name,
        local_account,
        remote_account,
        pair.num_rounds,
        pair.num_equal_stake,
        pair.total_stake,
        pair.total_diff,
        pair.remark,
    ]
    _write_data_row(ws, row, values)

    # Fill color by match type
    fill = _INTERNAL_FILL if pair.match_type == "内对打" else _EXTERNAL_FILL
    for col in range(1, len(values) + 1):
        ws.cell(row=row, column=col).fill = fill

    return row + 1


def _write_headers(ws, headers: list[str]) -> None:
    """Viết dòng header."""
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _THIN_BORDER


def _write_data_row(ws, row: int, values: list) -> None:
    """Viết 1 dòng data."""
    for col, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.font = _DATA_FONT
        cell.border = _THIN_BORDER

        if isinstance(val, str):
            # Account / text — force text format to preserve leading zeros
            cell.value = str(val)
            cell.data_type = 's'
            cell.alignment = _TEXT_ALIGN
            cell.number_format = '@'
        elif isinstance(val, float):
            cell.number_format = _MONEY_FORMAT
            cell.alignment = _DATA_ALIGN
        elif isinstance(val, int):
            cell.alignment = _DATA_ALIGN


def _auto_column_width(ws) -> None:
    """Auto-fit column width."""
    for col_cells in ws.columns:
        max_length = 0
        column = col_cells[0].column_letter
        for cell in col_cells:
            try:
                val = str(cell.value) if cell.value else ""
                # Estimate: CJK chars count as ~2, others as 1
                length = sum(2 if ord(c) > 127 else 1 for c in val)
                if length > max_length:
                    max_length = length
            except Exception:
                pass
        adjusted = min(max_length + 4, 50)
        ws.column_dimensions[column].width = max(adjusted, 10)
